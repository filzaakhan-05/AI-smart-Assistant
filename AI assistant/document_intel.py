"""Feature 2: Document Intelligence (PDF / TXT / DOCX)."""
import io
import re
from collections import Counter
from math import log

from config import ask

STOPWORDS = {
    "the", "a", "an", "and", "or", "of", "to", "in", "on", "for", "is", "are", "was",
    "were", "be", "it", "this", "that", "with", "as", "at", "by", "from", "what", "which",
    "who", "how", "do", "does", "did", "me", "my", "i", "about", "can", "you", "please",
}


# ---------- Loading ----------
def load_document(name: str, data: bytes) -> str:
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("PDF is password-protected.")
        pages = [(p.extract_text() or "").strip() for p in reader.pages]
        if not any(pages):
            raise ValueError("No extractable text (scanned/image-only PDF?).")
        return "\n".join(f"[Page {i + 1}]\n{t}" for i, t in enumerate(pages))
    if ext == "docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif ext in ("txt", "md"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type: .{ext}")
    if not text.strip():
        raise ValueError("The file contains no text.")
    return text


# ---------- Chunking & retrieval ----------
def chunk_text(text: str, size: int = 1200, overlap: int = 200) -> list[str]:
    chunks, start = [], 0
    step = max(1, size - overlap)
    while start < len(text):
        chunks.append(text[start:start + size])
        start += step
    return chunks


def _tokens(s: str) -> list[str]:
    return re.findall(r"[a-z0-9]+", s.lower())


def rank_chunks(query: str, chunks: list[tuple[str, str]], k: int = 5):
    """chunks = [(source, text)]. TF-IDF style keyword ranking. Returns top-k."""
    q = {w for w in _tokens(query) if w not in STOPWORDS} or set(_tokens(query))
    if not q or not chunks:
        return []
    docs = [Counter(_tokens(t)) for _, t in chunks]
    n = len(docs)
    df = {w: sum(1 for d in docs if w in d) for w in q}
    scored = []
    for i, ((src, text), d) in enumerate(zip(chunks, docs)):
        score = sum(d[w] * log(1 + n / df[w]) for w in q if df[w])
        if score > 0:
            scored.append((score, i, src, text))
    scored.sort(key=lambda x: (-x[0], x[1]))
    return [(src, text) for _, _, src, text in scored[:k]]


class DocumentStore:
    def __init__(self):
        self.docs: dict[str, str] = {}

    def add(self, name: str, data: bytes):
        self.docs[name] = load_document(name, data)

    def all_chunks(self) -> list[tuple[str, str]]:
        return [(n, c) for n, t in self.docs.items() for c in chunk_text(t)]

    def search(self, query: str, k: int = 5):
        return rank_chunks(query, self.all_chunks(), k)


# ---------- AI features ----------
def _long_text_reduce(text: str, instruction: str, limit: int = 40000) -> str:
    """Map-reduce (repeated until it fits) so long documents fit the context window."""
    for _ in range(4):
        if len(text) <= limit:
            return text
        parts = [text[i:i + 15000] for i in range(0, len(text), 15000)]
        notes = [ask("You condense text faithfully.", f"{instruction}\n\n{p}", 800) for p in parts]
        text = "\n\n".join(notes)
    return text[:limit]


def summarize(text: str) -> str:
    body = _long_text_reduce(text, "Summarize the key points of this section.")
    return ask(
        "You are a precise document analyst. Use only the provided text.",
        f"Write a concise summary (one paragraph + 5 bullet points) of:\n\n{body}",
    )


def key_topics(text: str) -> str:
    body = _long_text_reduce(text, "List the main topics in this section.")
    return ask(
        "You are a precise document analyst. Use only the provided text.",
        f"Identify the 5-10 key topics, each with a one-line description:\n\n{body}",
    )


def extract_info(text: str) -> str:
    body = _long_text_reduce(text, "Note important dates, names, numbers, deadlines, and obligations.")
    return ask(
        "You are a precise document analyst. Use only the provided text.",
        "Extract important information as grouped bullets: people/organizations, dates, "
        f"amounts/numbers, deadlines, obligations/action items:\n\n{body}",
    )


def answer_question(store: DocumentStore, question: str) -> str:
    hits = store.search(question, k=6)
    if not hits:
        return "I couldn't find anything relevant to that in the uploaded document(s)."
    context = "\n\n".join(f"[Source: {s}]\n{t}" for s, t in hits)
    return ask(
        "Answer ONLY using the provided excerpts. If the answer isn't in them, say "
        "'The document does not contain this information.' Do not use outside knowledge. "
        "Mention the source file (and page if shown).",
        f"Excerpts:\n{context}\n\nQuestion: {question}",
    )
